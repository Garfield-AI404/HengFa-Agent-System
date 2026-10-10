"""D19 R6 Excel 解析与行级转换测试（6.4 步骤 3、6.3 规则 5、6.5、RV-28）。

RV-28 通过标准逐条：
- 能定位文件哈希、映射版本和原始行号（标准行与错误行均携带）；
- 6.3 规则 5：日期、数字或枚举转换失败列出行号和原因；
- 行级错误不影响其他合法行（错误行单独进入 row_errors，合法行继续）；
- 6.5：优先在内存中解析（BytesIO），不产生临时文件；上传文件不进
  长期记忆（解析函数只接受 bytes，无落盘路径）。

映射定义复用 D18 MappingVersionRecord/MappingColumnRecord 结构。
"""

from __future__ import annotations

import io
from datetime import datetime, timezone

import pytest
from openpyxl import Workbook, load_workbook

from administrative_agent.files.excel_parser import (
    ParsedResult,
    RowError,
    StandardRow,
    convert_cell,
    parse_excel_bytes,
)
from administrative_agent.files.hashing import sha256_bytes
from administrative_agent.files.mapping_service import (
    MappingColumnRecord,
    MappingVersionRecord,
)

HEADER = ("型号", "日期", "单重", "数量", "单价", "状态", "备注")


def make_mapping(
    columns,
    *,
    sheet_name="对账单",
    header_row=1,
    version=3,
    template_key="recon-metal",
) -> MappingVersionRecord:
    return MappingVersionRecord(
        version_id=10,
        template_key=template_key,
        version=version,
        sheet_name=sheet_name,
        header_row=header_row,
        status="published",
        source_file_hash="sha256-mapping",
        created_by="manager01",
        created_at=datetime(2026, 8, 23, tzinfo=timezone.utc),
        columns=tuple(columns),
    )


def metal_columns():
    return [
        MappingColumnRecord(excel_column="型号", target_field="model",
                            data_type="string", is_required=True),
        MappingColumnRecord(excel_column="日期", target_field="in_date",
                            data_type="date", is_required=True),
        MappingColumnRecord(excel_column="单重", target_field="weight_per_unit",
                            data_type="decimal", conversion_rule="trim"),
        MappingColumnRecord(excel_column="数量", target_field="quantity",
                            data_type="int"),
        MappingColumnRecord(excel_column="单价", target_field="unit_price",
                            data_type="decimal"),
        MappingColumnRecord(excel_column="状态", target_field="status",
                            data_type="enum",
                            enum_mapping={"完成": "completed", "进行中": "in_progress"}),
        MappingColumnRecord(excel_column="备注", target_field="notes",
                            data_type="string", default_value=""),
    ]


def build_xlsx(rows, header=HEADER, sheet_name="对账单") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(header)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---- RV-28：定位信息（文件哈希 / 映射版本 / 原始行号） ----


def test_valid_rows_keep_hash_version_and_original_row_numbers():
    """6.4 步骤 3：标准化行保留文件哈希、映射版本和原始 Excel 行号。"""
    data = build_xlsx([
        ("M-1", "2026-06-02", 3.00, 10, 12.50, "完成", "ok"),
        ("M-2", "2026/6/3", 4.00, 5, 15.00, "进行中", ""),
    ])
    result = parse_excel_bytes(data, make_mapping(metal_columns()))
    assert isinstance(result, ParsedResult)
    assert result.valid_count == 2
    assert result.error_count == 0
    assert result.source_file_hash == sha256_bytes(data)
    assert result.mapping_version == 3
    row1, row2 = result.rows
    assert row1.row_number == 2 and row2.row_number == 3  # 原始 Excel 行号
    assert row1.values["model"] == "M-1"
    # 日期多种格式统一规范化为 YYYY-MM-DD
    assert row2.values["in_date"] == "2026-06-03"
    # 数字保留展示格式（比较口径由 rules 层规范化，此处不改格式）
    assert row1.values["weight_per_unit"] == "3"  # Excel 数值 3.00 原样读取
    assert row1.values["quantity"] == "10"
    assert row1.values["unit_price"] == "12.5"
    # 枚举转换
    assert row1.values["status"] == "completed"
    assert row2.values["status"] == "in_progress"
    # 默认值兜底（备注空 → default_value "" 进入 values）
    assert row2.values["notes"] == ""


def test_header_row_offset_keeps_real_row_numbers():
    """header_row=2（首行为标题行）：数据行号仍是 Excel 原始行号。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "对账单"
    ws.append(["示例五金 2026 年 6 月对账单"])  # 标题行（非表头）
    ws.append(HEADER)
    ws.append(["M-1", "2026-06-02", 3.00, 10, 12.50, "完成", ""])
    buf = io.BytesIO()
    wb.save(buf)
    mapping = make_mapping(metal_columns(), header_row=2)
    result = parse_excel_bytes(buf.getvalue(), mapping)
    assert result.valid_count == 1
    assert result.rows[0].row_number == 3  # 标题行 1 + 表头 2 + 数据 3


# ---- 6.3 规则 5：转换失败列出行号和原因 ----


def test_conversion_failures_listed_with_row_column_and_reason():
    """日期/数字/枚举转换失败逐行列出行号和原因，不阻断其他合法行。"""
    data = build_xlsx([
        ("M-1", "2026-06-02", 3.00, 10, 12.50, "完成", ""),
        ("M-2", "2026-13-99", 3.00, 10, 12.50, "完成", ""),   # 非法日期
        ("M-3", "2026-06-03", "abc", 10, 12.50, "完成", ""),   # 非法数字（单重）
        ("M-4", "2026-06-04", 3.00, "x", 12.50, "完成", ""),   # 非法整数（数量）
        ("M-5", "2026-06-05", 3.00, 10, 12.50, "不存在状态", ""),  # 非法枚举
        ("M-6", "2026-06-06", 3.00, 10, 12.50, "完成", ""),
    ])
    result = parse_excel_bytes(data, make_mapping(metal_columns()))
    assert result.valid_count == 2  # 第 1、6 行合法
    assert {r.row_number for r in result.rows} == {2, 7}
    assert result.error_count == 4
    by_row: dict[int, RowError] = {e.row_number: e for e in result.errors}
    assert 3 in by_row and "日期转换失败" in by_row[3].reason
    assert by_row[3].excel_column == "日期" and by_row[3].target_field == "in_date"
    assert 4 in by_row and "数字转换失败" in by_row[4].reason
    assert by_row[4].excel_column == "单重"
    assert 5 in by_row and "不是整数" in by_row[5].reason
    assert 6 in by_row and "枚举转换失败" in by_row[6].reason
    # 错误行也保留定位信息（RV-28：能定位文件哈希、映射版本和原始行号）
    assert result.source_file_hash == sha256_bytes(data)
    assert result.mapping_version == 3


def test_required_empty_value_is_row_error():
    """必填列为空 → 行级错误（列名 + 原因），该行不进标准行。"""
    data = build_xlsx([
        ("M-1", "2026-06-02", 3.00, 10, 12.50, "完成", ""),
        ("M-2", "", 3.00, 10, 12.50, "完成", ""),  # 必填日期为空
    ])
    result = parse_excel_bytes(data, make_mapping(metal_columns()))
    assert result.valid_count == 1
    assert result.rows[0].row_number == 2
    assert result.error_count == 1
    err = result.errors[0]
    assert err.row_number == 3
    assert err.excel_column == "日期" and "为空" in err.reason


def test_required_column_missing_in_header_reported():
    """表头缺少必填列 → 行级错误（防御表达，不猜测匹配）。"""
    data = build_xlsx(
        [("M-1", 3.00, 10, 12.50, "完成", "")],
        header=("型号", "单重", "数量", "单价", "状态", "备注"),
    )
    result = parse_excel_bytes(data, make_mapping(metal_columns()))
    assert result.error_count >= 1
    assert any(
        e.excel_column == "日期" and "表头缺少必填列" in e.reason
        for e in result.errors
    )


def test_default_value_applies_to_empty_cell():
    """default_value 兜底空值并仍走转换（失败同样列行号原因）。"""
    cols = [
        MappingColumnRecord(excel_column="数量", target_field="quantity",
                            data_type="int", is_required=True, default_value="0"),
    ]
    # 行不能全空（空行跳过）；数量为空但行内有其他非空列 → 默认值兜底
    data = build_xlsx([("", "x")], header=("数量", "备注"))
    result = parse_excel_bytes(data, make_mapping(cols))
    assert result.valid_count == 1
    assert result.rows[0].values["quantity"] == "0"
    # 默认值无法转换 → 行错误（fail closed，不静默用坏默认值）
    cols_bad = [
        MappingColumnRecord(excel_column="数量", target_field="quantity",
                            data_type="int", default_value="abc"),
    ]
    result_bad = parse_excel_bytes(data, make_mapping(cols_bad))
    assert result_bad.valid_count == 0
    assert any("不是整数" in e.reason for e in result_bad.errors)


def test_conversion_rule_trim_and_unknown_rule_fail_closed():
    """conversion_rule=trim 去除首尾空白；未知规则 fail closed 不静默忽略。"""
    cols_trim = [
        MappingColumnRecord(excel_column="单价", target_field="unit_price",
                            data_type="decimal", conversion_rule="trim"),
    ]
    data = build_xlsx([("  12.50  ",)], header=("单价",))
    result = parse_excel_bytes(data, make_mapping(cols_trim))
    assert result.rows[0].values["unit_price"] == "12.50"
    # 未知 conversion_rule → 该列行级错误（不静默忽略规则）
    cols_unknown = [
        MappingColumnRecord(excel_column="单价", target_field="unit_price",
                            data_type="decimal", conversion_rule="upper"),
    ]
    result_unknown = parse_excel_bytes(data, make_mapping(cols_unknown))
    assert result_unknown.valid_count == 0
    assert any("未知 conversion_rule" in e.reason for e in result_unknown.errors)


def test_convert_cell_unknown_data_type_fail_closed():
    """未知 data_type → fail closed（错误文本，不静默返回原值）。"""
    value, err = convert_cell(
        data_type="float", cell="1.5", enum_mapping=None, conversion_rule=None
    )
    assert value is None and err is not None and "未知 data_type" in err


def test_bool_and_int_conversion():
    from administrative_agent.files.excel_parser import _BOOL_TEXT

    value, err = convert_cell(data_type="bool", cell="是", enum_mapping=None,
                              conversion_rule=None)
    assert value == "true" and err is None
    value, err = convert_cell(data_type="bool", cell="maybe", enum_mapping=None,
                              conversion_rule=None)
    assert value is None and "布尔转换失败" in err
    value, err = convert_cell(data_type="int", cell=3.0, enum_mapping=None,
                              conversion_rule=None)
    assert value == "3" and err is None
    assert _BOOL_TEXT["否"] == "false"


# ---- 6.5：内存解析，不产生临时文件 ----


def test_parse_uses_bytes_stream_not_temp_file(monkeypatch):
    """6.5：优先内存解析——断言 openpyxl 收到的是 BytesIO 流而非文件路径，
    解析全程不产生临时文件。"""
    captured: list[object] = []

    def spy_load_workbook(source, *args, **kwargs):
        captured.append(source)
        return load_workbook(source, *args, **kwargs)

    monkeypatch.setattr("administrative_agent.files.excel_parser.load_workbook",
                        spy_load_workbook)
    data = build_xlsx([("M-1", "2026-06-02", 3.00, 10, 12.50, "完成", "")])
    result = parse_excel_bytes(data, make_mapping(metal_columns()))
    assert result.valid_count == 1
    assert captured, "解析必须调用 openpyxl 读取"
    assert isinstance(captured[0], io.BytesIO)  # 内存流，非文件路径
