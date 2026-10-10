"""映射模板写工具（mapping_template）。

把原本只对 manager 管理 API 开放（web/routes_mappings.py）的映射模板
导入/发布/退役（D18 mapping_service）扩展为 Agent 写工具，支持对话场景
「先建模板 → 发布 → 对账」闭环：

- ``create_draft``：附件 Excel（file_id）→ 表头列名 + 列定义双重校验
  （validate_column_definitions / check_required_and_unknown_columns，
  纯校验不落库）→ 确认卡（列数/校验 warnings）→ run 时 import_draft
  落新 draft 版本（版本号递增，6.3 规则 6 不覆盖历史）；
- ``publish``：get_version 存在性校验 → 确认卡 → run 时 publish
  （服务层 MappingStateError/MappingPublishedExistsError 兜底 fail
  closed——同 key 已有 published 被拒，规则 2）；
- ``retire``：同 publish（服务层 MappingStateError 兜底，终态不可恢复）。

**required_scopes=frozenset() 的决策**：映射写是本地 Postgres 配置表操作
（excel_mapping_versions/excel_mapping_columns），不触 ERP，无
erp:query/orders:* 等 scope 可言；安全由 manager 角色（写工具分支强制）+
确认卡（payload_hash 变化即失效，7.3）+ 任务/审计留痕（D6 task 行 +
EVENT_MAPPING 事件，批次 3 路由接线）兜底。

policy_family=("R6.mapping",)：6.3 处理规则 7 条的规则包（本工具参数轮
注入正文，10.2 要求 7——确定性校验在服务层，规则只约束模型行为）。
"""

from __future__ import annotations

import hashlib
import io
import os
from dataclasses import replace
from datetime import datetime, timezone
from typing import Any, Literal

from pydantic import BaseModel, Field, model_validator
from openpyxl import load_workbook

from administrative_agent.agent.tools.base import (
    PreparedToolWrite,
    ToolContext,
    ToolPreparationError,
    ToolResultSummary,
    _require_wiring,
)
from administrative_agent.config import settings
from administrative_agent.domains.orders import (
    ConfirmationCard,
    card_payload_hash,
    localize_confirmation_fields,
)
from administrative_agent.files.mapping_service import (
    MappingError,
    check_required_and_unknown_columns,
    validate_column_definitions,
)


class MappingColumnSpec(BaseModel):
    """单列映射定义（与 D18 MappingColumnRecord 同构；create_draft 载荷）。"""

    excel_column: str = Field(..., description="Excel 表头列名（定位列位置）")
    target_field: str = Field(..., description="目标 ERP 字段（须在 ERP_FIELD_CATALOG 内）")
    data_type: str = Field(
        ..., description="数据类型：string/int/decimal/date/enum/bool 之一"
    )
    conversion_rule: str | None = Field(None, description="转换规则：trim/none（未知规则拒绝）")
    is_required: bool = Field(False, description="是否必填（空值且无默认值 → 行级错误）")
    enum_mapping: dict[str, str] | None = Field(
        None, description="data_type=enum 时 Excel 文本 → 标准值映射"
    )
    default_value: str | None = Field(None, description="空单元格默认值（仍走转换）")


class MappingTemplateArgs(BaseModel):
    """mapping_template 参数（三动作：导入草稿/发布/退役）。"""

    action: Literal["create_draft", "publish", "retire"] = Field(
        ..., description="动作：create_draft（导入新草稿）/publish（发布草稿）/retire（退役版本）"
    )
    template_key: str | None = Field(
        None, description="映射模板键（create_draft 必填）"
    )
    sheet_name: str = Field("Sheet1", description="Excel 工作表名（create_draft）")
    header_row: int = Field(1, ge=1, description="表头所在 Excel 行（1 起始；create_draft）")
    file_id: str | None = Field(
        None, description="映射来源 Excel 附件 id（create_draft 必填）"
    )
    version_id: int | None = Field(
        None, description="映射版本 id（publish/retire 必填）"
    )
    columns: list[MappingColumnSpec] | None = Field(
        None, description="列映射定义（create_draft 必填，1 列以上）"
    )

    @model_validator(mode="after")
    def _validate_by_action(self) -> "MappingTemplateArgs":
        """按动作校验必填参数（fail closed：缺参不进入 prepare）。"""
        if self.action == "create_draft":
            if not self.template_key:
                raise ValueError("create_draft 必须提供 template_key")
            if not self.file_id:
                raise ValueError("create_draft 必须提供 file_id（映射来源 Excel 附件）")
            if not self.columns:
                raise ValueError("create_draft 必须提供 columns（列映射定义，至少 1 列）")
        elif self.version_id is None:
            raise ValueError(f"{self.action} 必须提供 version_id")
        return self


def prepare_mapping_template(
    ctx: ToolContext, args: MappingTemplateArgs
) -> PreparedToolWrite:
    """映射写确认前准备：只读校验 + 确认卡（不落库）。

    fail closed：依赖未接线/附件无效/列定义校验错误/目标版本不存在 →
    ToolPreparationError（不产生任何版本）。
    """
    _require_wiring(ctx, tool="mapping_template")
    if ctx.mapping_service is None:
        raise ToolPreparationError(
            "工具 mapping_template 未接线映射服务（harness 运行时缺 mapping_service），不能执行"
        )
    if ctx.context is None:
        raise ToolPreparationError("缺少会话上下文（context 未接线），不能执行")
    if args.action == "create_draft":
        return _prepare_create_draft(ctx, args)
    return _prepare_publish_retire(ctx, args)


def run_mapping_template(
    ctx: ToolContext, args: Any = None, *, task_id: str, execution: dict[str, Any]
) -> ToolResultSummary:
    """映射写执行：按批准的 execution 载荷执行 create_draft/publish/retire。

    fail closed：缺执行载荷/未知 action/服务层异常 → ok=False + 中文错误
    （不产生任何版本）。
    """
    _require_wiring(ctx, tool="mapping_template")
    err_prefix = {"tool": "mapping_template", "task_id": task_id}
    if ctx.mapping_service is None:
        return ToolResultSummary(
            ok=False,
            summary={
                **err_prefix,
                "error": "映射服务未接线（mapping_service），不能执行（fail closed）",
            },
        )
    if ctx.context is None:
        return ToolResultSummary(
            ok=False,
            summary={
                **err_prefix,
                "error": "缺少会话上下文（context 未接线），不能执行（fail closed）",
            },
        )
    action = execution.get("action")
    try:
        if action == "create_draft":
            return _run_create_draft(ctx, task_id, execution)
        if action == "publish":
            record = ctx.mapping_service.publish(
                actor=ctx.context.user_id, version_id=execution["version_id"]
            )
            return ToolResultSummary(
                ok=True,
                summary={
                    "action": "publish",
                    "template_key": record.template_key,
                    "version_id": record.version_id,
                    "version": record.version,
                    "status": record.status,
                },
            )
        if action == "retire":
            record = ctx.mapping_service.retire(
                actor=ctx.context.user_id, version_id=execution["version_id"]
            )
            return ToolResultSummary(
                ok=True,
                summary={
                    "action": "retire",
                    "template_key": record.template_key,
                    "version_id": record.version_id,
                    "version": record.version,
                    "status": record.status,
                },
            )
    except MappingError as exc:
        return ToolResultSummary(
            ok=False, summary={**err_prefix, "error": f"映射操作失败：{exc}"}
        )
    except Exception as exc:
        return ToolResultSummary(
            ok=False, summary={**err_prefix, "error": f"映射操作执行失败：{exc}"}
        )
    return ToolResultSummary(
        ok=False,
        summary={
            **err_prefix,
            "error": f"执行上下文缺少合法 action（{action!r}），不执行（fail closed）",
        },
    )


# ---------------------------------------------------------------------------
# create_draft：附件校验 + 列定义校验（纯校验不落库）
# ---------------------------------------------------------------------------


def _prepare_create_draft(
    ctx: ToolContext, args: MappingTemplateArgs
) -> PreparedToolWrite:
    if ctx.attachment_store is None:
        raise ToolPreparationError(
            "create_draft 需要附件仓储（harness 运行时缺 attachment_store），不能执行"
        )
    record = ctx.attachment_store.get(args.file_id)
    if record is None or _expired(record):
        raise ToolPreparationError(
            f"附件 {args.file_id!r} 不存在或已过期（fail closed），需重新上传映射文件"
        )
    if record.thread_id != ctx.context.thread_id:
        raise ToolPreparationError(
            f"附件 {args.file_id!r} 属于其他会话线程，拒绝使用（跨线程附件隔离，fail closed）"
        )
    data = _read_attachment_bytes(record)
    source_hash = hashlib.sha256(data).hexdigest()

    # 列定义合法性校验（6.3 规则 3/5；errors → fail closed 不产生版本）
    column_dicts = [c.model_dump() for c in (args.columns or [])]
    col_errors, col_warnings = validate_column_definitions(column_dicts)
    if col_errors:
        raise ToolPreparationError(
            "映射列定义校验失败（fail closed）：" + "；".join(col_errors)
        )
    # Excel 表头列名 + 必填列/未知列校验（规则 3/4；未知列仅告警展示）
    try:
        source_columns = _read_excel_header(data, args.sheet_name, args.header_row)
    except ToolPreparationError:
        raise
    req_errors, unknown_warnings = check_required_and_unknown_columns(
        columns=column_dicts,
        source_columns=source_columns,
        required_columns=None,
    )
    if req_errors:
        raise ToolPreparationError(
            "映射导入校验失败（fail closed）：" + "；".join(req_errors)
        )
    warnings = tuple(col_warnings) + tuple(unknown_warnings)

    # [2026-09-10 修复] hash_fields 显式声明「哈希用原始载荷」（无隐藏主键、
    # 无派生展示字段）。注意：本地化会改键名（template_key→模板键、
    # source_file_hash→源文件哈希，见 CONFIRMATION_FIELD_LABELS），故本卡
    # hash 与改造前口径不同；但本类卡执行期不复算 hash——is_confirmation_valid
    # 仅在 domains/orders.py 改单执行器与 domains/finance.py 收支改单执行器
    # 调用，确认环节比对的是 checkpoint 内同一份存储值，故无功能影响。
    raw_fields = (
        ("action", args.action),
        ("template_key", args.template_key),
        ("sheet_name", args.sheet_name),
        ("header_row", args.header_row),
        ("列数", len(column_dicts)),
        ("source_file_hash", source_hash),
        ("warnings", "；".join(warnings) if warnings else "无"),
    )
    card = ConfirmationCard(
        operation_type="映射模板导入（draft）",
        target=f"映射模板 {args.template_key}",
        fields_to_write=localize_confirmation_fields(raw_fields),
        hash_fields=raw_fields,
        auto_filled=(),
        expected_impact=(
            "在本地映射配置表新建 draft 版本"
            "（excel_mapping_versions/excel_mapping_columns）；不触 ERP。"
        ),
        limitations=(
            "draft 版本不参与对账（6.3 规则 1：只使用 published）；",
            "发布前校验：同 template_key 已有 published 版本则发布被拒"
            "（规则 2，须先退役当前版本）；",
            "Excel 中未映射列自动忽略并告警，不自行建立映射（规则 4）。",
        ),
        payload_hash="",
    )
    card = replace(card, payload_hash=card_payload_hash(card))
    return PreparedToolWrite(
        card=card.to_dict(),
        payload_hash=card.payload_hash,
        execution={
            "tool": "mapping_template",
            "action": args.action,
            "template_key": args.template_key,
            "version_id": args.version_id,
            "file_id": args.file_id,
            "sheet_name": args.sheet_name,
            "header_row": args.header_row,
            "columns": column_dicts,
            "source_file_hash": source_hash,
            "source_columns": list(source_columns),
        },
        policy_key="R6.mapping",
    )


def _run_create_draft(
    ctx: ToolContext, task_id: str, execution: dict[str, Any]
) -> ToolResultSummary:
    err_prefix = {"tool": "mapping_template", "task_id": task_id}
    file_id = execution.get("file_id")
    template_key = execution.get("template_key")
    if not file_id or not template_key:
        return ToolResultSummary(
            ok=False,
            summary={
                **err_prefix,
                "error": "执行上下文缺少 file_id/template_key（fail closed，不执行）",
            },
        )
    if ctx.attachment_store is None:
        return ToolResultSummary(
            ok=False,
            summary={
                **err_prefix,
                "error": "附件仓储未接线（attachment_store），不能执行（fail closed）",
            },
        )
    record = ctx.attachment_store.get(file_id)
    if record is None or _expired(record):
        return ToolResultSummary(
            ok=False,
            summary={
                **err_prefix,
                "error": f"附件 {file_id!r} 不存在或已过期，拒绝导入（fail closed）",
            },
        )
    try:
        data = _read_attachment_bytes(record)
    except ToolPreparationError as exc:
        return ToolResultSummary(ok=False, summary={**err_prefix, "error": str(exc)})
    # 重读字节计算实际哈希（import_draft 落库值以实际文件为准）
    source_hash = hashlib.sha256(data).hexdigest()
    result = ctx.mapping_service.import_draft(
        actor=ctx.context.user_id,
        template_key=template_key,
        sheet_name=execution.get("sheet_name") or "Sheet1",
        header_row=int(execution.get("header_row") or 1),
        source_file_hash=source_hash,
        columns=execution.get("columns") or [],
        source_columns=execution.get("source_columns"),
    )
    created = result.version
    return ToolResultSummary(
        ok=True,
        summary={
            "action": "create_draft",
            "template_key": created.template_key,
            "version_id": created.version_id,
            "version": created.version,
            "status": created.status,
            "source_file_hash": source_hash,
            "warnings": list(result.warnings),
        },
    )


# ---------------------------------------------------------------------------
# publish/retire：目标版本存在性校验（服务层状态机兜底 fail closed）
# ---------------------------------------------------------------------------


def _prepare_publish_retire(
    ctx: ToolContext, args: MappingTemplateArgs
) -> PreparedToolWrite:
    try:
        record = ctx.mapping_service.get_version(args.version_id)
    except MappingError as exc:
        raise ToolPreparationError(
            f"映射版本 {args.version_id} 不存在（fail closed）：{exc}"
        ) from exc
    action_label = "映射模板发布" if args.action == "publish" else "映射模板退役"
    # [2026-09-10 修复] hash_fields：与上方 draft 卡同理——显式用原始载荷哈希。
    # 本地化改键名（template_key→模板键）故 hash 与改造前口径不同，但本类卡
    # 执行期不复算 hash（is_confirmation_valid 仅在 orders/finance 改单执行器
    # 调用），确认比对的是 checkpoint 内同一份存储值，无功能影响。
    raw_fields = (
        ("action", args.action),
        ("version_id", args.version_id),
        ("template_key", record.template_key),
        ("version", record.version),
        ("status", record.status),
        ("列数", len(record.columns)),
    )
    card = ConfirmationCard(
        operation_type=action_label,
        target=f"映射版本 v{record.version}（{record.template_key}，"
        f"当前状态 {record.status}）",
        fields_to_write=localize_confirmation_fields(raw_fields),
        hash_fields=raw_fields,
        auto_filled=(),
        expected_impact=(
            f"{record.template_key} 映射版本 v{record.version} 状态变更"
            "（draft→published→retired 状态机）；不触 ERP。"
        ),
        limitations=(
            "发布前校验：同 template_key 已有 published 版本则发布被拒"
            "（规则 2，一次只有一个发布版本）；",
            "非 draft 状态发布/退役已退役版本 → 服务层拒绝（fail closed）；",
            "退役（retired）为终态，历史不可恢复（6.3 状态机）。",
        ),
        payload_hash="",
    )
    card = replace(card, payload_hash=card_payload_hash(card))
    return PreparedToolWrite(
        card=card.to_dict(),
        payload_hash=card.payload_hash,
        execution={
            "tool": "mapping_template",
            "action": args.action,
            "template_key": args.template_key or record.template_key,
            "version_id": args.version_id,
            "file_id": args.file_id,
        },
        policy_key="R6.mapping",
    )


# ---------------------------------------------------------------------------
# 内部辅助
# ---------------------------------------------------------------------------


def _expired(record: Any) -> bool:
    """附件是否已过期（expires_at 非空且 <= now，与仓储清理口径一致）。"""
    expires_at = getattr(record, "expires_at", None)
    if expires_at is None:
        return False
    return expires_at <= datetime.now(timezone.utc)


def _read_attachment_bytes(record: Any) -> bytes:
    """附件内容读取（storage_path 为 attachment_dir 下的相对路径；
    读取失败 fail closed）。"""
    path = os.path.join(settings.attachment_dir, record.storage_path or "")
    try:
        with open(path, "rb") as f:
            return f.read()
    except OSError as exc:
        raise ToolPreparationError(
            f"附件文件读取失败（{path}）：{exc}（fail closed）"
        ) from exc


def _read_excel_header(data: bytes, sheet_name: str, header_row: int) -> list[str]:
    """读取附件 Excel 表头行单元格文本（源列名清单，供规则 3/4 校验）。"""
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ToolPreparationError(
            f"附件 Excel 解析失败：{exc}（fail closed）"
        ) from exc
    try:
        sheet = wb[sheet_name] if sheet_name in wb.sheetnames else wb.active
        header: tuple[Any, ...] | None = None
        for i, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            if i == header_row:
                header = row
                break
    finally:
        wb.close()
    if header is None:
        raise ToolPreparationError(
            f"附件 Excel 未找到表头行（header_row={header_row}，fail closed）"
        )
    return [str(c).strip() for c in header if c is not None and str(c).strip()]
