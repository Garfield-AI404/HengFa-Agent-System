"""Excel 解析与行级转换（规格 2.6 files/excel_parser.py；6.4 步骤 3、
6.3 规则 5、6.5 文件保留）。

- 按 published 映射（D18 mapping_service 读取）解析客户订单 Excel 成
  标准化行：header_row 定位表头、列级转换规则应用（6.3 规则 5：
  日期、数字或枚举转换失败时列出行号和原因）；
- 保留定位信息：文件哈希（sha256，6.4 步骤 3）、映射版本、原始行号
  （RV-28 通过标准：能定位文件哈希、映射版本和原始行号）；
- 行级转换错误不静默、不影响其他合法行（错误行单独进入 row_errors，
  合法行继续）；
- 6.5：优先在内存中解析（openpyxl read_only 从 BytesIO 读取），
  不使用临时文件——上传文件不进入长期记忆，checkpoint/审计只保存
  标准化数据、文件哈希、映射版本和行号，不保存无关工作表内容。
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any, Mapping, Sequence

from openpyxl import load_workbook

from administrative_agent.files.hashing import sha256_bytes
from administrative_agent.files.mapping_service import MappingVersionRecord

#: 合法 data_type（与 D18 mapping_service.DATA_TYPES 一致；转换规则 5）
VALID_TYPES = ("string", "int", "decimal", "date", "enum", "bool")

#: 已知 conversion_rule（列级转换规则；未知规则 fail closed 不静默忽略）
KNOWN_RULES = frozenset({"trim", "none"})

#: 日期可接受的文本格式（YYYY-MM-DD / YYYY/M/D / YYYY.M.D / YYYY年M月D日）
_DATE_PATTERNS = (
    re.compile(r"^\s*(\d{4})-(\d{1,2})-(\d{1,2})\s*$"),
    re.compile(r"^\s*(\d{4})/(\d{1,2})/(\d{1,2})\s*$"),
    re.compile(r"^\s*(\d{4})\.(\d{1,2})\.(\d{1,2})\s*$"),
    re.compile(r"^\s*(\d{4})年(\d{1,2})月(\d{1,2})日\s*$"),
)

#: bool 文本集合（Excel 单元格常见布尔表达）→ 标准值
_BOOL_TEXT = {
    "true": "true", "false": "false",
    "是": "true", "否": "false",
    "1": "true", "0": "false",
    "yes": "true", "no": "false",
}


@dataclass(frozen=True)
class RowError:
    """行级转换错误（6.3 规则 5：列出行号和原因）。

    RV-28 定位信息由调用方与报告关联：source_file_hash/mapping_version
    记录在 ParsedResult 与差异报告中，row_number/excel_column 逐条给出。
    """

    row_number: int  # Excel 原始行号
    excel_column: str  # 出错的 Excel 列名
    target_field: str
    reason: str  # 转换失败原因（规则 5"列出行号和原因"）


@dataclass(frozen=True)
class StandardRow:
    """一行 Excel 数据的标准化结果（6.5：checkpoint 保存标准化数据）。"""

    row_number: int  # Excel 原始行号
    values: Mapping[str, str]  # target_field -> 转换后展示值（字符串）
    source_file_hash: str
    mapping_version: int


@dataclass(frozen=True)
class ParsedResult:
    """一次 Excel 解析结果（合法行 + 行级错误；RV-28 定位信息）。"""

    rows: tuple[StandardRow, ...]
    errors: tuple[RowError, ...]
    source_file_hash: str
    mapping_version: int

    @property
    def valid_count(self) -> int:
        return len(self.rows)

    @property
    def error_count(self) -> int:
        return len(self.errors)


# ---------------------------------------------------------------------------
# 单元格转换（6.3 规则 5：失败返回原因文本，不抛异常——错误逐行表达）
# ---------------------------------------------------------------------------


def _text(cell: Any) -> str:
    """单元格 → 展示文本（None → 空串；数字/日期转字符串保持原格式）。"""
    if cell is None:
        return ""
    if isinstance(cell, datetime):
        return cell.strftime("%Y-%m-%d")
    if isinstance(cell, date):
        return cell.strftime("%Y-%m-%d")
    if isinstance(cell, float) and cell.is_integer():
        return str(int(cell))
    return str(cell)


def _convert_date(cell: Any) -> str | None:
    """日期转换 → 'YYYY-MM-DD'；失败返回 None（规则 5 由调用方组织原因）。"""
    if isinstance(cell, datetime):
        return cell.strftime("%Y-%m-%d")
    if isinstance(cell, date):
        return cell.strftime("%Y-%m-%d")
    text = _text(cell).strip()
    if not text:
        return None
    for pat in _DATE_PATTERNS:
        m = pat.match(text)
        if m:
            y, mo, d = (int(g) for g in m.groups())
            try:
                return date(y, mo, d).strftime("%Y-%m-%d")
            except ValueError:
                return None
    return None


def _convert_decimal(cell: Any) -> str | None:
    """十进制转换 → 规范化展示（保留输入文本，去首尾空白）；失败 None。"""
    try:
        Decimal(_text(cell).strip())
    except InvalidOperation:
        return None
    return _text(cell).strip()


def _convert_int(cell: Any) -> str | None:
    try:
        value = int(_text(cell).strip())
    except ValueError:
        return None
    return str(value)


def _convert_bool(cell: Any) -> str | None:
    key = _text(cell).strip().lower()
    return _BOOL_TEXT.get(key)


def convert_cell(
    *,
    data_type: str,
    cell: Any,
    enum_mapping: Mapping[str, str] | None,
    conversion_rule: str | None,
) -> tuple[str | None, str | None]:
    """按映射列定义转换一个单元格。

    返回 (value, error)：value=None 且 error 非空 → 转换失败（规则 5）；
    conversion_rule 未知 → fail closed（不静默忽略规则）。
    """
    if conversion_rule and conversion_rule not in KNOWN_RULES:
        return None, f"未知 conversion_rule={conversion_rule!r}（不静默忽略）"
    if conversion_rule == "trim":
        raw = _text(cell).strip()
        cell = raw
    if data_type == "string":
        if conversion_rule == "trim":
            return _text(cell).strip(), None
        return _text(cell), None
    if data_type == "int":
        value = _convert_int(cell)
        if value is None:
            return None, f"数字转换失败：{_text(cell)!r} 不是整数"
        return value, None
    if data_type == "decimal":
        value = _convert_decimal(cell)
        if value is None:
            return None, f"数字转换失败：{_text(cell)!r} 不是有效数字"
        return value, None
    if data_type == "date":
        value = _convert_date(cell)
        if value is None:
            return None, f"日期转换失败：{_text(cell)!r} 无法解析为日期"
        return value, None
    if data_type == "bool":
        value = _convert_bool(cell)
        if value is None:
            return None, f"布尔转换失败：{_text(cell)!r} 不是 true/false/是/否"
        return value, None
    if data_type == "enum":
        key = _text(cell).strip()
        if key not in (enum_mapping or {}):
            return None, f"枚举转换失败：{key!r} 不在映射 {sorted((enum_mapping or {}))} 中"
        return enum_mapping[key], None  # type: ignore[index]
    return None, f"未知 data_type={data_type!r}（不静默忽略）"


# ---------------------------------------------------------------------------
# 解析入口
# ---------------------------------------------------------------------------


def parse_excel_bytes(
    data: bytes,
    mapping: MappingVersionRecord,
) -> ParsedResult:
    """内存解析客户订单 Excel（6.5：不落盘，不进入长期记忆）。

    - header_row：表头所在 Excel 行（映射定义），表头单元格文本与
      excel_column 匹配定位列；
    - 数据行 = header_row 之后的非空行；行号保留 Excel 原始行号；
    - 逐列按 data_type 应用转换（规则 5），失败记录 RowError（行号+列名
      +原因），不阻断其他行；
    - 必填列（is_required）空值且无默认值 → 行错误；
    - default_value：空单元格使用（仍走转换，失败记录错误）。
    """
    source_hash = sha256_bytes(data)
    version = mapping.version
    if not mapping.columns:
        return ParsedResult((), (), source_hash, version)

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        return _parse_workbook(wb, mapping, source_hash, version)
    finally:
        wb.close()


def _parse_workbook(
    wb: Any, mapping: MappingVersionRecord, source_hash: str, version: int
) -> ParsedResult:
    sheet = wb[mapping.sheet_name] if mapping.sheet_name in wb.sheetnames else wb.active
    rows_iter = sheet.iter_rows(values_only=True)

    # header_row（1 起始）：跳过其前行，读取表头行
    header: tuple[Any, ...] | None = None
    for i, row in enumerate(rows_iter, start=1):
        if i == mapping.header_row:
            header = row
            break
    if header is None:
        return ParsedResult((), (), source_hash, version)

    # 表头文本 → excel_column 定位（映射列在表头中的列位置）
    col_index: dict[str, int] = {}  # excel_column -> 0 起始列号
    for idx, cell in enumerate(header):
        name = str(cell).strip() if cell is not None else ""
        if name:
            col_index.setdefault(name, idx)

    parsed_rows: list[StandardRow] = []
    errors: list[RowError] = []
    for row_number, row in enumerate(rows_iter, start=mapping.header_row + 1):
        if all(c is None or str(c).strip() == "" for c in row):
            continue  # 空行跳过
        values: dict[str, str] = {}
        row_bad = False
        for col_def in mapping.columns:
            excel_column = col_def.excel_column
            idx = col_index.get(excel_column)
            if idx is None:
                # 表头缺该列：必填列 → 行错误；非必填列跳过（6.3 规则 3
                # 已在导入阶段保证，此处防御）
                if col_def.is_required:
                    errors.append(RowError(
                        row_number, excel_column, col_def.target_field,
                        f"表头缺少必填列 {excel_column!r}",
                    ))
                    row_bad = True
                continue
            cell = row[idx] if idx < len(row) else None
            if cell is None or (isinstance(cell, str) and cell.strip() == ""):
                if col_def.default_value is not None:
                    cell = col_def.default_value
                elif col_def.is_required:
                    errors.append(RowError(
                        row_number, excel_column, col_def.target_field,
                        f"必填列 {excel_column!r} 为空",
                    ))
                    row_bad = True
                    continue
                else:
                    continue  # 非必填空值不转换、不进 values
            value, err = convert_cell(
                data_type=col_def.data_type,
                cell=cell,
                enum_mapping=col_def.enum_mapping,
                conversion_rule=col_def.conversion_rule,
            )
            if err is not None:
                errors.append(RowError(
                    row_number, excel_column, col_def.target_field, err,
                ))
                row_bad = True
                continue
            assert value is not None
            values[col_def.target_field] = value
        if row_bad:
            continue  # 有转换错误的行不进入标准行（错误已逐条记录）
        if values:
            parsed_rows.append(StandardRow(
                row_number=row_number,
                values=dict(values),
                source_file_hash=source_hash,
                mapping_version=version,
            ))
    return ParsedResult(
        rows=tuple(parsed_rows), errors=tuple(errors),
        source_file_hash=source_hash, mapping_version=version,
    )
